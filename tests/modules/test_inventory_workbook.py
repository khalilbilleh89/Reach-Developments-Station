"""The operator's inventory workbook: four sheets, one transaction, no formulas.

The CSV importer this sits beside is more capable and is not going anywhere.
What it cannot do is let a phase exist before a unit mentions it, so a team
laying out a development has to invent units to describe a building. These
sheets can, which is the whole reason the workbook exists — and the reason most
of the tests below are about hierarchy that has no units under it yet.

Two properties carry the rest. **A workbook is data**: a cell containing
``=1+1`` is refused, never resolved, and the value Excel cached is never read.
**A workbook is one batch**: four sheets validate together and apply together,
so a unit failing on the last row takes that workbook's phases with it.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.audit.models import AuditEvent
from app.modules.inventory import workbook
from app.modules.inventory.models import Building, Floor, Phase, Unit
from tests.modules.conftest import inventory_url

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --------------------------------------------------------------------------- #
# Building workbooks to send
# --------------------------------------------------------------------------- #


def _book(
    *,
    phases: tuple[dict[str, object], ...] = (),
    buildings: tuple[dict[str, object], ...] = (),
    floors: tuple[dict[str, object], ...] = (),
    units: tuple[dict[str, object], ...] = (),
    version: str | None = None,
    edit: Callable[[Workbook], None] | None = None,
) -> bytes:
    """A filled-in copy of the real template.

    Built from ``workbook.template()`` rather than from a hand-made file, so
    every test here also proves the template the operator downloads is the file
    the parser accepts. A fixture assembled independently could drift from the
    template and none of these tests would notice.
    """
    book = load_workbook(io.BytesIO(workbook.template()))
    if version is not None:
        book[workbook.SHEET_INSTRUCTIONS][workbook.VERSION_CELL] = version
    for name, rows in (
        (workbook.SHEET_PHASES, phases),
        (workbook.SHEET_BUILDINGS, buildings),
        (workbook.SHEET_FLOORS, floors),
        (workbook.SHEET_UNITS, units),
    ):
        sheet = book[name]
        columns = [column for column, _ in workbook.SHEET_COLUMNS[name]]
        for offset, row in enumerate(rows, start=workbook.FIRST_DATA_ROW):
            for position, column in enumerate(columns, start=1):
                if column in row:
                    sheet.cell(row=offset, column=position, value=row[column])
    if edit is not None:
        edit(book)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


PHASE = {"phase_code": "PHASE-1", "phase_name": "Phase 1"}
BUILDING = {"phase_code": "PHASE-1", "building_code": "B1", "building_name": "Building 1"}
FLOOR = {
    "phase_code": "PHASE-1",
    "building_code": "B1",
    "floor_code": "01",
    "floor_label": "First floor",
}


def _unit(number: str, reference: str, **overrides: object) -> dict[str, object]:
    row = {
        "phase_code": "PHASE-1",
        "building_code": "B1",
        "floor_code": "01",
        "unit_number": number,
        "unit_reference": reference,
        "asset_class": "apartment",
    }
    row.update(overrides)
    return row


def _validate(client: TestClient, project_id: str, body: bytes, mode: str = "create") -> dict:
    response = client.post(
        f"{inventory_url(project_id)}/import/workbook/validate?mode={mode}",
        content=body,
        headers={"Content-Type": XLSX},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _apply(client: TestClient, project_id: str, body: bytes, mode: str = "create") -> dict:
    response = client.post(
        f"{inventory_url(project_id)}/import/workbook/apply?mode={mode}",
        content=body,
        headers={"Content-Type": XLSX},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _messages(report: dict) -> str:
    return " | ".join(
        f"{issue['sheet']}:{issue['row']}:{issue['column']}:{issue['message']}"
        for issue in report["issues"]
    )


# --------------------------------------------------------------------------- #
# The template itself
# --------------------------------------------------------------------------- #


class TestTheTemplateIsTheContract:
    """What an operator downloads is what the parser reads.

    Every assertion here is about the file rather than about a constant, because
    a template generated from the same tuple it is checked against proves only
    that the tuple exists.
    """

    def test_the_download_is_a_real_workbook_named_for_the_project(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        response = admin_client.get(f"{inventory_url(operational_project)}/import/template.xlsx")

        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(XLSX)
        assert "inventory-template.xlsx" in response.headers["content-disposition"]
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            assert "xl/workbook.xml" in archive.namelist()

    def test_it_carries_every_sheet_the_parser_requires(self) -> None:
        book = load_workbook(io.BytesIO(workbook.template()))

        assert book.sheetnames == list(workbook.REQUIRED_SHEETS)

    def test_the_template_version_is_in_a_cell_not_the_filename(self) -> None:
        """A downloaded file is renamed, copied and mailed on. A cell is not."""
        book = load_workbook(io.BytesIO(workbook.template()))
        notes = book[workbook.SHEET_INSTRUCTIONS]

        assert notes[workbook.VERSION_LABEL_CELL].value == workbook.VERSION_LABEL
        assert notes[workbook.VERSION_CELL].value == workbook.TEMPLATE_VERSION
        assert workbook.TEMPLATE_VERSION in workbook.SUPPORTED_VERSIONS

    def test_every_sheet_has_exactly_the_headers_the_contract_names(self) -> None:
        book = load_workbook(io.BytesIO(workbook.template()))

        for name, columns in workbook.SHEET_COLUMNS.items():
            written = [cell.value for cell in book[name][workbook.HEADER_ROW]]
            assert written == [column for column, _ in columns], name

    def test_the_units_sheet_speaks_only_columns_the_importer_already_knows(self) -> None:
        """One vocabulary, not two.

        The Units sheet is a subset of the CSV importer's columns on purpose: a
        column here that the shared row reader had never heard of would be a
        second definition of what a unit is.
        """
        from app.modules.inventory import import_service

        for column, _ in workbook.UNIT_COLUMNS:
            assert column in import_service.KNOWN_COLUMNS, column

    def test_it_carries_no_formula_anywhere(self) -> None:
        """The template must not teach the shape the parser refuses."""
        book = load_workbook(io.BytesIO(workbook.template()))

        for name in book.sheetnames:
            for row in book[name].iter_rows():
                for cell in row:
                    assert not (isinstance(cell.value, str) and cell.value.startswith("=")), (
                        f"{name}!{cell.coordinate}"
                    )


# --------------------------------------------------------------------------- #
# Hierarchy that exists before any unit does
# --------------------------------------------------------------------------- #


class TestStructureCanExistBeforeUnits:
    """The reason this workbook exists at all.

    The CSV importer materialises a phase because some unit row named it, so a
    development cannot be laid out before its units are known. Every test here
    loads hierarchy with no unit under it.
    """

    def test_a_workbook_of_phases_alone_applies(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _apply(
            admin_client,
            operational_project,
            _book(phases=(PHASE, {"phase_code": "PHASE-2", "phase_name": "Phase 2"})),
        )

        assert report["applied"] is True, _messages(report)
        assert report["structure"]["records"]["phases"] == {"create": 2, "update": 0}
        assert report["structure"]["records"]["units"] == {"create": 0, "update": 0}
        codes = set(db.scalars(select(Phase.code)))
        assert {"PHASE-1", "PHASE-2"} <= codes

    def test_a_phase_and_a_building_with_no_units_applies(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _apply(
            admin_client, operational_project, _book(phases=(PHASE,), buildings=(BUILDING,))
        )

        assert report["applied"] is True, _messages(report)
        assert db.scalars(select(Building.code)).all() == ["B1"]

    def test_a_floor_with_no_units_applies(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _apply(
            admin_client,
            operational_project,
            _book(phases=(PHASE,), buildings=(BUILDING,), floors=(FLOOR,)),
        )

        assert report["applied"] is True, _messages(report)
        assert db.scalars(select(Floor.code)).all() == ["01"]
        assert db.scalars(select(Unit.id)).all() == []

    def test_the_whole_hierarchy_and_its_units_apply_as_one(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _apply(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(FLOOR,),
                units=(_unit("101", "B1-101"), _unit("102", "B1-102")),
            ),
        )

        assert report["applied"] is True, _messages(report)
        assert report["structure"]["records"]["units"] == {"create": 2, "update": 0}
        assert report["create_count"] == 5
        assert sorted(db.scalars(select(Unit.unit_reference))) == ["B1-101", "B1-102"]

    def test_a_phase_sheet_needs_the_permission_the_phase_form_needs(
        self, engineer_client: TestClient, operational_project: str
    ) -> None:
        """A workbook is another way to ask, never a way to ask for more.

        Creating a phase through the form requires a project configurer.
        Design and engineering may write structure but not phases, so the same
        person meets the same refusal here — before any row is read.
        """
        response = engineer_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(phases=(PHASE,)),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 403, response.text

    def test_the_same_person_may_still_load_buildings_and_floors(
        self, engineer_client: TestClient, admin_client: TestClient, operational_project: str
    ) -> None:
        """The refusal above is about phases, not about workbooks."""
        _apply(admin_client, operational_project, _book(phases=(PHASE,)))

        report = _apply(
            engineer_client, operational_project, _book(buildings=(BUILDING,), floors=(FLOOR,))
        )

        assert report["applied"] is True, _messages(report)


# --------------------------------------------------------------------------- #
# Parents, duplicates and the messages an operator has to act on
# --------------------------------------------------------------------------- #


class TestTheReportNamesTheCellToFix:
    """A validation report is read by somebody holding the spreadsheet.

    "Foreign key violation" tells them nothing. Sheet, row, column and a
    sentence about the actual problem tells them where to click.
    """

    def test_a_building_under_an_undeclared_phase_names_the_sheet_and_the_phase(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(admin_client, operational_project, _book(buildings=(BUILDING,)))

        assert report["error_count"] == 1
        issue = report["issues"][0]
        assert issue["sheet"] == workbook.SHEET_BUILDINGS
        assert issue["row"] == 2
        assert issue["column"] == "phase_code"
        assert "PHASE-1" in issue["message"]
        assert workbook.SHEET_PHASES in issue["message"]

    def test_a_unit_on_an_undeclared_floor_is_refused_rather_than_inventing_one(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        """The one thing tolerating missing hierarchy must not buy.

        The unit reader is told to accept hierarchy that does not exist yet,
        because most of it arrives on the sheets above. Without this check a
        typo in floor_code would quietly bring a nameless floor into being.
        """
        report = _validate(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(FLOOR,),
                units=(_unit("101", "B1-101", floor_code="99"),),
            ),
        )

        assert report["error_count"] >= 1, _messages(report)
        assert any(
            issue["sheet"] == workbook.SHEET_UNITS and issue["column"] == "floor_code"
            for issue in report["issues"]
        ), _messages(report)
        assert db.scalars(select(Floor.code)).all() == []

    def test_the_same_phase_twice_is_a_duplicate_not_a_merge(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(
            admin_client,
            operational_project,
            _book(phases=(PHASE, {"phase_code": "PHASE-1", "phase_name": "Phase One"})),
        )

        assert report["error_count"] == 1
        assert report["issues"][0]["sheet"] == workbook.SHEET_PHASES
        assert "row 2" in report["issues"][0]["message"]

    def test_an_identical_duplicate_is_still_a_duplicate(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        """Silently keeping one of two identical rows hides the mistake.

        An operator who pasted a block twice needs to know they did.
        """
        report = _validate(admin_client, operational_project, _book(phases=(PHASE, dict(PHASE))))

        assert report["error_count"] == 1

    def test_two_buildings_with_one_code_in_one_phase_collide(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(
            admin_client,
            operational_project,
            _book(phases=(PHASE,), buildings=(BUILDING, {**BUILDING, "building_name": "Other"})),
        )

        assert report["error_count"] == 1
        assert report["issues"][0]["column"] == "building_code"

    def test_the_same_building_code_in_two_phases_is_two_buildings(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        """B1 in Phase 1 and B1 in Phase 2 are different buildings.

        The domain scopes a building code to its phase, so refusing this would
        be the importer inventing a rule the rest of the system does not have.
        """
        report = _apply(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE, {"phase_code": "PHASE-2", "phase_name": "Phase 2"}),
                buildings=(BUILDING, {**BUILDING, "phase_code": "PHASE-2"}),
            ),
        )

        assert report["applied"] is True, _messages(report)
        assert sorted(db.scalars(select(Building.code))) == ["B1", "B1"]

    def test_two_floors_with_one_code_in_one_building_collide(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(FLOOR, {**FLOOR, "floor_label": "Other"}),
            ),
        )

        assert report["error_count"] == 1
        assert report["issues"][0]["column"] == "floor_code"

    def test_the_same_floor_code_in_two_buildings_is_two_floors(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _apply(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(
                    BUILDING,
                    {**BUILDING, "building_code": "B2", "building_name": "Building 2"},
                ),
                floors=(FLOOR, {**FLOOR, "building_code": "B2"}),
            ),
        )

        assert report["applied"] is True, _messages(report)
        assert sorted(db.scalars(select(Floor.code))) == ["01", "01"]

    def test_a_floor_code_keeps_the_text_a_building_actually_uses(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        """B2, GF and RF are floor codes. An integer cannot hold them."""
        report = _apply(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(
                    {**FLOOR, "floor_code": "B1", "floor_label": "Basement 1", "level_number": -1},
                    {**FLOOR, "floor_code": "GF", "floor_label": "Ground floor", "level_number": 0},
                    {**FLOOR, "floor_code": "RF", "floor_label": "Roof", "level_number": 9},
                ),
            ),
        )

        assert report["applied"] is True, _messages(report)
        assert sorted(db.scalars(select(Floor.code))) == ["B1", "GF", "RF"]


# --------------------------------------------------------------------------- #
# A spreadsheet is data
# --------------------------------------------------------------------------- #


class TestNothingInAWorkbookIsExecuted:
    """The property that makes accepting spreadsheets acceptable.

    An inventory import is data capture. A workbook that arrives with formulas
    in it is refused rather than resolved, and the value the last version of
    Excel happened to cache is never read — that number was computed by
    something nobody here has seen, from inputs nobody here can check.
    """

    def test_a_formula_in_a_data_cell_refuses_the_workbook(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        def put_a_formula(book: Workbook) -> None:
            book[workbook.SHEET_PHASES]["B2"] = "=1+1"

        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(phases=(PHASE,), edit=put_a_formula),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text
        assert "formula" in response.json()["detail"].lower()
        # And emphatically not the number two.
        assert response.json()["detail"] != "2"

    def test_the_cached_value_of_a_formula_is_never_read(self) -> None:
        """Proved on the loader rather than through the API.

        ``data_only=True`` would hand back whatever Excel last calculated, and
        the parse would import it as though somebody had typed it. This asserts
        the parser sees the formula text — which is what lets it refuse.
        """
        body = _book(
            phases=(PHASE,), edit=lambda book: book[workbook.SHEET_PHASES].__setitem__("B2", "=1+1")
        )
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=False)
        try:
            assert workbook._is_formula(book[workbook.SHEET_PHASES]["B2"])
        finally:
            book.close()

    def test_a_formula_in_a_header_refuses_the_workbook(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(edit=lambda book: book[workbook.SHEET_UNITS].__setitem__("A1", "=A2")),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text

    def test_a_macro_enabled_workbook_is_refused(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        body = _book(phases=(PHASE,))
        with_macro = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(body)) as source,
            zipfile.ZipFile(with_macro, "w") as target,
        ):
            for entry in source.infolist():
                target.writestr(entry, source.read(entry.filename))
            target.writestr("xl/vbaProject.bin", b"\x00\x01\x02")

        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=with_macro.getvalue(),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text
        assert "macro" in response.json()["detail"].lower()

    def test_a_file_that_is_not_a_workbook_is_refused(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=b"phase_code,phase_name\nPHASE-1,Phase 1\n",
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text

    def test_a_password_protected_workbook_says_so(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        """An encrypted .xlsx is an OLE2 container, not a ZIP.

        Left to the ZIP reader it produces "not a readable workbook", which
        sends an operator to rebuild a file that is fine.
        """
        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512,
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text
        assert "password" in response.json()["detail"].lower()

    def test_a_file_over_the_bound_is_refused_before_it_is_parsed(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=b"\x00" * (workbook.MAX_BYTES + 1),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text

    def test_the_uncompressed_bound_is_the_one_that_matters(self) -> None:
        """A workbook is a ZIP, so the size on the wire is not the size in memory.

        The compressed bound is the operator's number. This one is the defence:
        a few hundred kilobytes that unpack to gigabytes is refused on the
        archive listing, before a byte is decompressed.
        """
        assert workbook.MAX_UNCOMPRESSED > workbook.MAX_BYTES
        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/workbook.xml", b"\x00" * (workbook.MAX_UNCOMPRESSED + 1))

        with pytest.raises(workbook.WorkbookRefused):
            workbook._archive_is_safe(bomb.getvalue())


class TestTheTemplateVersionIsAContract:
    def test_an_unknown_version_is_refused_by_name(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(phases=(PHASE,), version="inventory_structure_v99"),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text
        detail = response.json()["detail"]
        assert "inventory_structure_v99" in detail
        assert workbook.TEMPLATE_VERSION in detail

    def test_a_blank_version_is_refused(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(phases=(PHASE,), version=""),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text

    def test_the_accepted_version_is_reported_back(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(admin_client, operational_project, _book(phases=(PHASE,)))

        assert report["template_version"] == workbook.TEMPLATE_VERSION

    def test_a_renamed_sheet_is_refused(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        def rename(book: Workbook) -> None:
            book[workbook.SHEET_FLOORS].title = "Levels"

        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(edit=rename),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text
        assert workbook.SHEET_FLOORS in response.json()["detail"]

    def test_a_renamed_column_is_refused_rather_than_guessed_at(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        def rename(book: Workbook) -> None:
            book[workbook.SHEET_PHASES]["A1"] = "phase"

        response = admin_client.post(
            f"{inventory_url(operational_project)}/import/workbook/validate?mode=create",
            content=_book(edit=rename),
            headers={"Content-Type": XLSX},
        )

        assert response.status_code == 422, response.text
        assert "phase_code" in response.json()["detail"]


# --------------------------------------------------------------------------- #
# Fields, modes, and the transaction
# --------------------------------------------------------------------------- #


class TestFieldsAreJudgedByTheDomain:
    def test_an_asset_class_the_domain_does_not_have_is_refused(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(FLOOR,),
                units=(_unit("101", "B1-101", asset_class="spaceship"),),
            ),
        )

        assert report["error_count"] >= 1
        assert any(issue["sheet"] == workbook.SHEET_UNITS for issue in report["issues"])

    def test_a_phase_status_the_domain_does_not_have_is_refused_by_name(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(
            admin_client, operational_project, _book(phases=({**PHASE, "status": "nearly"},))
        )

        assert report["error_count"] == 1
        assert report["issues"][0]["column"] == "status"
        assert "planning" in report["issues"][0]["message"]

    def test_a_date_in_another_locale_is_refused_rather_than_guessed(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        """03/04/2026 is April in one country and March in another.

        Picking one would silently move a phase's planned start by a month.
        """
        report = _validate(
            admin_client,
            operational_project,
            _book(phases=({**PHASE, "planned_start": "03/04/2026"},)),
        )

        assert report["error_count"] == 1
        assert report["issues"][0]["column"] == "planned_start"
        assert workbook.DATE_FORMAT in report["issues"][0]["message"]

    def test_an_iso_date_is_stored(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _apply(
            admin_client,
            operational_project,
            _book(phases=({**PHASE, "planned_start": "2026-04-20", "status": "active"},)),
        )

        assert report["applied"] is True, _messages(report)
        phase = db.scalars(select(Phase).where(Phase.code == "PHASE-1")).one()
        assert phase.planned_start.isoformat() == "2026-04-20"
        assert phase.status == "active"

    def test_a_real_date_cell_is_read_as_a_date_not_a_timestamp(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        from datetime import datetime

        report = _apply(
            admin_client,
            operational_project,
            _book(phases=({**PHASE, "planned_start": datetime(2026, 4, 20, 0, 0)},)),
        )

        assert report["applied"] is True, _messages(report)
        phase = db.scalars(select(Phase).where(Phase.code == "PHASE-1")).one()
        assert phase.planned_start.isoformat() == "2026-04-20"

    def test_a_blank_optional_cell_writes_nothing(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        _apply(admin_client, operational_project, _book(phases=(PHASE,)))

        phase = db.scalars(select(Phase).where(Phase.code == "PHASE-1")).one()
        assert phase.notes is None
        assert phase.planned_start is None

    def test_a_missing_required_cell_names_the_column(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        report = _validate(
            admin_client, operational_project, _book(phases=({"phase_code": "PHASE-1"},))
        )

        assert report["error_count"] == 1
        assert report["issues"][0]["column"] == "phase_name"


class TestCreateAndUpsert:
    def test_create_mode_refuses_a_code_that_already_exists(
        self, admin_client: TestClient, operational_project: str
    ) -> None:
        _apply(admin_client, operational_project, _book(phases=(PHASE,)))

        report = _validate(admin_client, operational_project, _book(phases=(PHASE,)), mode="create")

        assert report["error_count"] == 1
        assert "create mode" in report["issues"][0]["message"]

    def test_upsert_updates_what_the_code_names(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        _apply(admin_client, operational_project, _book(phases=(PHASE,)))

        report = _apply(
            admin_client,
            operational_project,
            _book(phases=({**PHASE, "phase_name": "Phase One, renamed", "status": "active"},)),
            mode="upsert",
        )

        assert report["applied"] is True, _messages(report)
        assert report["structure"]["records"]["phases"] == {"create": 0, "update": 1}
        phase = db.scalars(select(Phase).where(Phase.code == "PHASE-1")).one()
        assert phase.name == "Phase One, renamed"
        assert phase.status == "active"
        assert db.scalars(select(Phase.id)).all() == [phase.id]

    def test_upsert_leaves_the_code_alone(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        """A code is identity. An import updates what a record says, not which record it is."""
        _apply(admin_client, operational_project, _book(phases=(PHASE,)))
        before = db.scalars(select(Phase).where(Phase.code == "PHASE-1")).one().id

        _apply(
            admin_client,
            operational_project,
            _book(phases=({**PHASE, "phase_name": "Renamed"},)),
            mode="upsert",
        )

        assert db.scalars(select(Phase).where(Phase.code == "PHASE-1")).one().id == before


class TestOneWorkbookIsOneTransaction:
    def test_validation_writes_nothing(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        report = _validate(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(FLOOR,),
                units=(_unit("101", "B1-101"),),
            ),
        )

        assert report["applied"] is False
        assert report["error_count"] == 0, _messages(report)
        assert db.scalars(select(Phase.id)).all() == []
        assert db.scalars(select(Unit.id)).all() == []

    def test_apply_reads_the_bytes_again_and_does_not_trust_an_earlier_validate(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        """A clean validate is not a token that unlocks apply.

        The file on the operator's disk can change between the two calls, and a
        caller can send any summary it likes.
        """
        clean = _book(phases=(PHASE,))
        assert _validate(admin_client, operational_project, clean)["error_count"] == 0

        broken = _book(phases=(PHASE, dict(PHASE)))
        report = _apply(admin_client, operational_project, broken)

        assert report["applied"] is False
        assert report["error_count"] == 1
        assert db.scalars(select(Phase.id)).all() == []

    def test_a_failure_while_writing_units_takes_the_whole_workbook_with_it(
        self,
        admin_client: TestClient,
        operational_project: str,
        db: Session,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """The interesting failure is the late one.

        Phases, buildings and floors are written before the first unit. If the
        unit write then fails, every one of them has to go — a development
        holding three floors and none of its units is a catalogue nobody can
        reason about, and it looks exactly like a successful import of a
        smaller file.
        """
        from app.modules.inventory import import_service

        def fail_late(*_args: object, **_kwargs: object) -> None:
            raise RuntimeError("row 247 of 247")

        monkeypatch.setattr(import_service, "apply_rows", fail_late)

        # Re-raised rather than swallowed: a batch that failed must look like a
        # batch that failed. A rollback that returned quietly would leave the
        # operator reading a success and believing three floors exist.
        with pytest.raises(RuntimeError, match="row 247"):
            admin_client.post(
                f"{inventory_url(operational_project)}/import/workbook/apply?mode=create",
                content=_book(
                    phases=(PHASE,),
                    buildings=(BUILDING,),
                    floors=(FLOOR,),
                    units=(_unit("101", "B1-101"),),
                ),
                headers={"Content-Type": XLSX},
            )

        assert db.scalars(select(Phase.id)).all() == []
        assert db.scalars(select(Building.id)).all() == []
        assert db.scalars(select(Floor.id)).all() == []
        assert db.scalars(select(Unit.id)).all() == []

    def test_one_workbook_writes_one_import_event_carrying_its_correlation(
        self, admin_client: TestClient, operational_project: str, db: Session
    ) -> None:
        _apply(
            admin_client,
            operational_project,
            _book(
                phases=(PHASE,),
                buildings=(BUILDING,),
                floors=(FLOOR,),
                units=(_unit("101", "B1-101"),),
            ),
        )

        events = db.scalars(
            select(AuditEvent).where(AuditEvent.action == "inventory.import_applied")
        ).all()
        assert len(events) == 1
        assert events[0].after_data["format"] == "workbook"
        assert events[0].after_data["template_version"] == workbook.TEMPLATE_VERSION
        assert events[0].correlation_id is not None
