"""FAQ template roundtrips, atomic writes, permissions and retained removal."""

import io

import pytest
from fastapi.testclient import TestClient
from httpx2 import Response
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from tests.factories import client_for, make_user
from tests.modules.conftest import grant_access


def root(project_id: str) -> str:
    return f"/api/v1/projects/{project_id}/faqs"


def workbook(client: TestClient, project_id: str, records: list[tuple[str | None, ...]]) -> bytes:
    response = client.get(f"{root(project_id)}/import/template")
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book["FAQs"]
    for number, row in enumerate(records, 2):
        for column, value in enumerate(row, 1):
            sheet.cell(number, column, value)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def send(client: TestClient, project_id: str, body: bytes, operation: str = "apply") -> Response:
    return client.post(f"{root(project_id)}/import/{operation}", content=body)


def test_template_validate_apply_replay_delete(
    sales_ops_client: TestClient, project_id: str, db: Session
) -> None:
    client = sales_ops_client
    answer = "First line\nنرحب بكم — ready to share"
    body = workbook(
        client, project_id, [("Visit?", answer), (None, None), ("Contact?", "Call the team")]
    )
    report = send(client, project_id, body, "validate").json()
    assert report["total_rows"] == 2 and report["created_count"] == 0
    assert report["error_count"] == 0 and not report["applied"]
    assert client.get(root(project_id)).json()["items"] == []
    assert (
        db.scalar(text("SELECT count(*) FROM audit_events WHERE entity_type='commercial_faq'")) == 0
    )
    report = send(client, project_id, body).json()
    assert report["applied"] and report["created_count"] == 2
    rows = client.get(root(project_id)).json()["items"]
    visit = next(row for row in rows if row["question"] == "Visit?")
    assert visit["answer"] == answer
    repeated = send(client, project_id, body).json()
    assert not repeated["applied"] and repeated["error_count"] == 2
    path = f"{root(project_id)}/{visit['id']}"
    assert client.delete(path, params={"version": 1, "reason": "Replace sample"}).status_code == 204
    assert client.delete(path, params={"version": 1, "reason": "Again"}).status_code == 404
    db.rollback()
    event = db.execute(
        text("SELECT before_data FROM audit_events WHERE action='faq.deleted'")
    ).scalar_one()
    assert event["answer"] == answer
    assert len(client.get(root(project_id)).json()["items"]) == 1


@pytest.mark.parametrize(
    "records",
    [
        [("Good", "Answer"), ("Missing answer", "")],
        [("Question", "Answer"), (" question ", "Duplicate")],
        [("Good", "Answer"), ("Long answer", "x" * 20001)],
        [("x" * 501, "Answer")],
    ],
    ids=["blank", "duplicate", "answer-limit", "question-limit"],
)
def test_invalid_batch_creates_nothing(
    sales_ops_client: TestClient,
    project_id: str,
    db: Session,
    records: list[tuple[str | None, ...]],
) -> None:
    body = workbook(sales_ops_client, project_id, records)
    result = send(sales_ops_client, project_id, body)
    assert result.status_code == 200, result.text
    report = result.json()
    assert not report["applied"] and report["created_count"] == 0
    assert report["issues"][0]["row"] >= 2
    assert sales_ops_client.get(root(project_id)).json()["items"] == []
    assert (
        db.scalar(text("SELECT count(*) FROM audit_events WHERE entity_type='commercial_faq'")) == 0
    )


@pytest.mark.parametrize(
    "case", ["formula", "header", "project", "empty", "extra-column", "too-many"]
)
def test_rejects_invalid_workbooks(
    sales_ops_client: TestClient, project_id: str, case: str
) -> None:
    body = workbook(sales_ops_client, project_id, [("Question", "Answer")])
    book = load_workbook(io.BytesIO(body))
    sheet = book["FAQs"]
    if case == "formula":
        sheet["B2"] = "=1+1"
    elif case == "header":
        sheet["A1"] = "Answer"
    elif case == "project":
        book["Instructions"]["B2"] = "another-project"
    elif case == "empty":
        sheet["A2"] = sheet["B2"] = None
    elif case == "extra-column":
        sheet["C2"] = "Do not lose me"
    else:
        for number in range(3, 1003):
            sheet.cell(number, 1, f"Question {number}")
            sheet.cell(number, 2, "Answer")
    out = io.BytesIO()
    book.save(out)
    book.close()
    assert send(sales_ops_client, project_id, out.getvalue()).status_code == 422
    assert sales_ops_client.get(root(project_id)).json()["items"] == []


def test_apply_rechecks_after_validation(sales_ops_client: TestClient, project_id: str) -> None:
    client = sales_ops_client
    body = workbook(client, project_id, [("Question", "Imported answer")])
    assert send(client, project_id, body, "validate").json()["error_count"] == 0
    assert (
        client.post(
            root(project_id), json={"question": "QUESTION", "answer": "Manual answer"}
        ).status_code
        == 201
    )
    result = send(client, project_id, body).json()
    assert result["error_count"] == 1 and not result["applied"]
    assert client.get(root(project_id)).json()["items"][0]["answer"] == "Manual answer"


def test_permissions_before_file_parsing(
    advisor_client: TestClient,
    engineer_client: TestClient,
    admin_client: TestClient,
    project_id: str,
    db: Session,
) -> None:
    for client in [advisor_client, engineer_client]:
        assert client.get(f"{root(project_id)}/import/template").status_code == 403
        for operation in ["validate", "apply"]:
            assert send(client, project_id, b"not a workbook", operation).status_code == 403
    restricted = make_user(db, email="faq-import-phase@example.com", roles=("sales_operations",))
    grant_access(admin_client, project_id, restricted)
    admin_client.patch(
        f"/api/v1/projects/{project_id}/access/{restricted.id}/phase-scope",
        json={"phase_scope": "selected"},
    )
    with client_for(restricted.email) as client:
        assert send(client, project_id, b"invalid").status_code == 403
    outsider = make_user(db, email="faq-import-outsider@example.com", roles=("sales_operations",))
    with client_for(outsider.email) as client:
        assert send(client, project_id, b"invalid").status_code == 404


def test_bounded_upload_and_malformed_package(
    sales_ops_client: TestClient, project_id: str
) -> None:
    assert send(sales_ops_client, project_id, b"x" * (5 * 1024 * 1024 + 1)).status_code == 413
    assert send(sales_ops_client, project_id, b"not a zip").status_code == 422
