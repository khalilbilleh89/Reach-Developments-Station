"""Bulk registration preserves manual permissions, ownership and atomicity."""

import io
import uuid
import zipfile

import pytest
from fastapi.testclient import TestClient
from httpx2 import Response
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.xlsx import MAX_BYTES
from tests.modules.conftest import sales_url


def workbook(client: TestClient, project: str, kind: str, records: list[dict]) -> bytes:
    response = client.get(f"{sales_url(project)}/import/{kind}/template")
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book[kind.title()]
    columns = [cell.value for cell in sheet[1]]
    for number, record in enumerate(records, 2):
        for column, key in enumerate(columns, 1):
            sheet.cell(number, column, record.get(key))
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def upload(
    client: TestClient, project: str, kind: str, body: bytes, operation: str = "apply"
) -> Response:
    return client.post(
        f"{sales_url(project)}/import/{kind}/{operation}",
        content=body,
        headers={"Content-Type": "application/octet-stream"},
    )


def test_agents_validate_import_retry_and_delete(
    sales_ops_client: TestClient, admin_client: TestClient, operational_project: str, db: Session
) -> None:
    base = sales_url(operational_project)
    body = workbook(
        sales_ops_client,
        operational_project,
        "agents",
        [
            {
                "display_name": "Agent One",
                "country": "Jordan",
                "branch": "Amman",
                "branch_leader": "Leader",
            },
            {"display_name": "Agent Two"},
        ],
    )
    checked = upload(sales_ops_client, operational_project, "agents", body, "validate")
    assert checked.status_code == 200, checked.text
    assert checked.json()["error_count"] == 0
    assert checked.json()["applied"] is False
    assert db.scalar(text("SELECT count(*) FROM sales_agents")) == 0
    assert db.scalar(text("SELECT count(*) FROM audit_events WHERE entity_type='sales_agent'")) == 0
    db.rollback()
    applied = upload(sales_ops_client, operational_project, "agents", body)
    assert applied.status_code == 200, applied.text
    assert applied.json()["created_count"] == 2
    agents = sales_ops_client.get(base + "/agents").json()
    assert len(agents) == 2
    assert agents[0]["branch_leader"] == "Leader"
    retried = upload(sales_ops_client, operational_project, "agents", body)
    assert retried.json()["applied"] is True
    assert retried.json()["created_count"] == 0
    assert retried.json()["duplicate_count"] == 2
    assert retried.json()["error_count"] == 0
    row = agents[1]
    assert (
        sales_ops_client.delete(
            f"{base}/agents/{row['id']}", params={"reason": "Import correction"}
        ).status_code
        == 204
    )
    assert (
        sales_ops_client.delete(
            f"{base}/agents/{row['id']}", params={"reason": "Repeated"}
        ).status_code
        == 404
    )
    assert sales_ops_client.get(f"{base}/agents/{row['id']}").status_code == 404


def test_buyers_keep_contacts_agent_and_sole_or_joint_ownership(
    advisor_client: TestClient,
    sales_ops_client: TestClient,
    finance_client: TestClient,
    admin_client: TestClient,
    operational_project: str,
    db: Session,
) -> None:
    base = sales_url(operational_project)
    agent = sales_ops_client.post(
        base + "/agents", json={"display_name": "Agent A", "country": "Jordan", "branch": "Amman"}
    ).json()
    template = advisor_client.get(base + "/import/buyers/template")
    book = load_workbook(io.BytesIO(template.content))
    assert book["Agent reference"]["A2"].value == agent["id"]
    assert book["Buyers"]["C2"].number_format == "@"
    book.close()
    body = workbook(
        advisor_client,
        operational_project,
        "buyers",
        [
            {
                "display_name": "Sole Buyer",
                "email": "sole@example.test",
                "phone": "0012345",
                "agent_id": agent["id"],
            },
            {"display_name": "Joint Buyers", "sole_purchaser": False},
        ],
    )
    checked = upload(advisor_client, operational_project, "buyers", body, "validate")
    assert checked.status_code == 200, checked.text
    assert checked.json()["error_count"] == 0, checked.text
    assert db.scalar(text("SELECT count(*) FROM clients")) == 0
    assert db.scalar(text("SELECT count(*) FROM client_parties")) == 0
    db.rollback()
    result = upload(advisor_client, operational_project, "buyers", body)
    assert result.json()["applied"] is True, result.text
    buyers = advisor_client.get(base + "/clients").json()
    sole = next(row for row in buyers if row["display_name"] == "Sole Buyer")
    joint = next(row for row in buyers if row["display_name"] == "Joint Buyers")
    assert sole["phone"] == "0012345"
    assert sole["agent_id"] == agent["id"]
    assert sales_ops_client.get(f"{base}/agents/{agent['id']}").json()["country"] == "Jordan"
    assert sole["owner_advisor_user_id"] is not None
    parties = advisor_client.get(f"{base}/clients/{sole['id']}/parties").json()
    assert len(parties) == 1 and parties[0]["share_fraction"] == "1.000000"
    assert advisor_client.get(f"{base}/clients/{joint['id']}/parties").json() == []
    assert "email" not in finance_client.get(f"{base}/clients/{sole['id']}").json()
    retried = upload(advisor_client, operational_project, "buyers", body).json()
    assert retried["applied"] is True
    assert retried["created_count"] == 0
    assert retried["duplicate_count"] == 2
    assert (
        sales_ops_client.delete(
            f"{base}/agents/{agent['id']}", params={"reason": "Referenced"}
        ).status_code
        == 409
    )
    assert (
        admin_client.delete(
            f"{base}/clients/{sole['id']}", params={"reason": "Correct import"}
        ).status_code
        == 204
    )
    assert (
        admin_client.delete(f"{base}/clients/{sole['id']}", params={"reason": "Again"}).status_code
        == 404
    )
    assert db.scalar(text("SELECT count(*) FROM audit_events WHERE action='client.deleted'")) == 1


@pytest.mark.parametrize("kind", ["buyers", "agents"])
def test_one_invalid_row_rolls_back_the_entire_batch(
    sales_ops_client: TestClient, operational_project: str, db: Session, kind: str
) -> None:
    records = [
        {"display_name": "Valid first"},
        {
            "display_name": " ",
            **(
                {"email": "missing-name@example.test"}
                if kind == "buyers"
                else {"country": "Jordan"}
            ),
        },
    ]
    body = workbook(sales_ops_client, operational_project, kind, records)
    for operation in ("validate", "apply"):
        response = upload(sales_ops_client, operational_project, kind, body, operation)
        assert response.status_code == 200, response.text
        assert response.json()["applied"] is False
        assert response.json()["issues"][0]["row"] == 3
        assert response.json()["issues"][0]["column"] == "display_name"
        table = "clients" if kind == "buyers" else "sales_agents"
        assert db.scalar(text(f"SELECT count(*) FROM {table}")) == 0
        db.rollback()


@pytest.mark.parametrize(
    "bad", ["formula", "headers", "empty", "project", "version", "macro", "not_xlsx"]
)
def test_bad_workbooks_are_refused_without_writes(
    sales_ops_client: TestClient, operational_project: str, db: Session, bad: str
) -> None:
    records = [{"display_name": "One"}]
    if bad == "formula":
        records[0]["display_name"] = "=1+1"
    body = workbook(sales_ops_client, operational_project, "agents", records)
    book = load_workbook(io.BytesIO(body))
    if bad == "headers":
        book["Agents"]["B1"] = "display_name"
    if bad == "project":
        book["Instructions"]["B2"] = str(uuid.uuid4())
    if bad == "version":
        book["Instructions"]["B1"] = "sales_buyers_v1"
    if bad == "empty":
        book["Agents"]["A2"] = None
    output = io.BytesIO()
    book.save(output)
    book.close()
    body = output.getvalue()
    if bad == "macro":
        output = io.BytesIO(body)
        with zipfile.ZipFile(output, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"synthetic")
        body = output.getvalue()
    if bad == "not_xlsx":
        body = b"not a workbook"
    response = upload(sales_ops_client, operational_project, "agents", body)
    assert response.status_code == 422, response.text
    assert db.scalar(text("SELECT count(*) FROM sales_agents")) == 0


def test_agent_assignment_rechecked_on_apply(
    sales_ops_client: TestClient, operational_project: str, db: Session
) -> None:
    base = sales_url(operational_project)
    agent = sales_ops_client.post(
        base + "/agents", json={"display_name": "Unavailable later"}
    ).json()
    body = workbook(
        sales_ops_client,
        operational_project,
        "buyers",
        [{"display_name": "First"}, {"display_name": "Second", "agent_id": agent["id"]}],
    )
    assert (
        upload(sales_ops_client, operational_project, "buyers", body, "validate").json()[
            "error_count"
        ]
        == 0
    )
    assert (
        sales_ops_client.patch(
            f"{base}/agents/{agent['id']}", json={"is_active": False}
        ).status_code
        == 200
    )
    response = upload(sales_ops_client, operational_project, "buyers", body)
    assert response.json()["applied"] is False
    assert response.json()["issues"][0]["row"] == 3
    assert db.scalar(text("SELECT count(*) FROM clients")) == 0


def test_import_authorization_and_project_scope(
    sales_ops_client: TestClient,
    finance_client: TestClient,
    advisor_client: TestClient,
    operational_project: str,
) -> None:
    base = sales_url(operational_project)
    body = workbook(sales_ops_client, operational_project, "buyers", [{"display_name": "One"}])
    for kind in ("buyers", "agents"):
        assert finance_client.get(f"{base}/import/{kind}/template").status_code == 403
        assert upload(finance_client, operational_project, kind, body).status_code == 403
        assert upload(sales_ops_client, str(uuid.uuid4()), kind, body).status_code == 404
    assert upload(advisor_client, operational_project, "buyers", body).status_code == 200
    assert (
        upload(sales_ops_client, operational_project, "buyers", b"x" * (MAX_BYTES + 1)).status_code
        == 413
    )


@pytest.mark.parametrize("kind", ["buyers", "agents"])
def test_expanding_list_skips_duplicates_and_rechecks_at_apply(
    sales_ops_client: TestClient, operational_project: str, db: Session, kind: str
) -> None:
    client = sales_ops_client
    project = operational_project
    base = sales_url(project)
    table = "clients" if kind == "buyers" else "sales_agents"
    records = [{"display_name": "First"}, {"display_name": "Second"}, {"display_name": " first "}]
    body = workbook(client, project, kind, records)
    checked = upload(client, project, kind, body, "validate").json()
    assert checked["new_count"] == 2 and checked["duplicate_count"] == 1
    assert checked["duplicates"] == [{"row": 4, "source": "file"}]
    assert checked["error_count"] == 0 and checked["created_count"] == 0
    assert db.scalar(text(f"SELECT count(*) FROM {table}")) == 0
    db.rollback()
    result = upload(client, project, kind, body).json()
    assert result["applied"] and result["created_count"] == 2
    assert result["duplicate_count"] == 1
    route = "clients" if kind == "buyers" else "agents"
    existing_row = client.get(f"{base}/{route}").json()[0]
    assert (
        client.patch(f"{base}/{route}/{existing_row['id']}", json={"is_active": False}).status_code
        == 200
    )
    audit_before = db.scalar(text("SELECT count(*) FROM audit_events"))
    parties_before = db.scalar(text("SELECT count(*) FROM client_parties"))
    db.rollback()
    retried = upload(client, project, kind, body).json()
    assert retried["applied"] and retried["created_count"] == 0
    assert retried["new_count"] == 0 and retried["duplicate_count"] == 3
    assert all(row["source"] == "existing" for row in retried["duplicates"])
    assert client.get(f"{base}/{route}/{existing_row['id']}").json()["is_active"] is False
    assert db.scalar(text("SELECT count(*) FROM audit_events")) == audit_before
    assert db.scalar(text("SELECT count(*) FROM client_parties")) == parties_before
    db.rollback()
    records.append({"display_name": "Third"})
    expanded = workbook(client, project, kind, records)
    checked = upload(client, project, kind, expanded, "validate").json()
    assert checked["new_count"] == 1 and checked["duplicate_count"] == 3
    route = "clients" if kind == "buyers" else "agents"
    # Another writer adds the new name after preview. Apply must skip it too.
    assert client.post(f"{base}/{route}", json={"display_name": "Third"}).status_code == 201
    applied = upload(client, project, kind, expanded).json()
    assert applied["applied"] and applied["created_count"] == 0
    assert applied["duplicate_count"] == 4
    records.append({"display_name": "Fourth"})
    final = upload(client, project, kind, workbook(client, project, kind, records)).json()
    assert final["applied"] and final["created_count"] == 1
    assert final["new_count"] == 1 and final["duplicate_count"] == 4
    assert db.scalar(text(f"SELECT count(*) FROM {table}")) == 4


@pytest.mark.parametrize("kind", ["buyers", "agents"])
def test_skipped_duplicates_do_not_hide_real_errors(
    sales_ops_client: TestClient, operational_project: str, db: Session, kind: str
) -> None:
    client = sales_ops_client
    project = operational_project
    first = workbook(client, project, kind, [{"display_name": "Existing"}])
    assert upload(client, project, kind, first).json()["created_count"] == 1
    invalid = {"email": "missing@example.test"} if kind == "buyers" else {"country": "Jordan"}
    body = workbook(
        client, project, kind, [{"display_name": "Existing"}, {"display_name": "New"}, invalid]
    )
    for operation in ("validate", "apply"):
        report = upload(client, project, kind, body, operation).json()
        assert report["applied"] is False and report["created_count"] == 0
        assert report["new_count"] == 1 and report["duplicate_count"] == 1
        assert report["error_count"] > 0 and report["issues"][0]["row"] == 4
        table = "clients" if kind == "buyers" else "sales_agents"
        assert db.scalar(text(f"SELECT count(*) FROM {table}")) == 1
        db.rollback()


@pytest.mark.parametrize("kind", ["buyers", "agents"])
def test_same_name_with_different_identity_fields_is_not_silently_skipped(
    sales_ops_client: TestClient, operational_project: str, kind: str
) -> None:
    distinguishing = "email" if kind == "buyers" else "branch"
    body = workbook(
        sales_ops_client,
        operational_project,
        kind,
        [
            {"display_name": "Same name", distinguishing: "first@example.test"},
            {"display_name": "Same name", distinguishing: "second@example.test"},
        ],
    )
    report = upload(sales_ops_client, operational_project, kind, body).json()
    assert report["applied"] and report["created_count"] == 2
    assert report["duplicate_count"] == 0


def test_duplicate_checks_do_not_disclose_another_advisors_buyers(
    sales_ops_client: TestClient, advisor_client: TestClient, operational_project: str, db: Session
) -> None:
    fields = {"display_name": "Scoped buyer", "email": "scoped@example.test"}
    base = sales_url(operational_project)
    assert sales_ops_client.post(base + "/clients", json=fields).status_code == 201
    body = workbook(advisor_client, operational_project, "buyers", [fields])
    checked = upload(advisor_client, operational_project, "buyers", body, "validate").json()
    assert checked["new_count"] == 1 and checked["duplicate_count"] == 0
    assert checked["duplicates"] == []
    saved = upload(advisor_client, operational_project, "buyers", body).json()
    assert saved["created_count"] == 1 and saved["duplicate_count"] == 0
    assert db.scalar(text("SELECT count(*) FROM clients")) == 2
