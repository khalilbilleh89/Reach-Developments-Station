"""Excel uses stable project/client/stage identity and preserves atomic updates."""

import io
import uuid
from datetime import date

import pytest
from fastapi.testclient import TestClient
from httpx2 import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from tests.modules.test_sales_operations import payload, read, stage_payload, url


def workbook(client: TestClient, project: str) -> Workbook:
    response = client.get(f"{url(project)}/import/template")
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content))


def upload(
    client: TestClient, project: str, book: Workbook, operation: str = "validate"
) -> Response:
    stream = io.BytesIO()
    book.save(stream)
    return client.post(
        f"{url(project)}/import/{operation}",
        params={"reason": "Verified signed originals"},
        content=stream.getvalue(),
    )


def test_template_prepopulates_clients_and_custom_stages(
    sales_ops_client: TestClient, project_id: str, buyer_id: str, db: Session
) -> None:
    data = read(sales_ops_client, project_id)
    draft = stage_payload(data)
    draft["stages"].append({"label": "Custom certificate", "section": "property_purchase"})
    assert sales_ops_client.put(f"{url(project_id)}/pipeline", json=draft).status_code == 204
    book = workbook(sales_ops_client, project_id)
    sheet = book["Clients"]
    assert sheet["A3"].value == buyer_id
    assert sheet["D3"].value == data["buyers"][0]["name"]
    assert "Custom certificate" in sheet.cell(2, 25).value
    assert sheet.freeze_panes == "G3"
    assert sheet.column_dimensions["A"].hidden
    assert len(sheet.data_validations.dataValidation) == 2
    report = upload(sales_ops_client, project_id, book).json()
    assert report["changed_clients"] == 0
    assert not report["issues"]
    assert db.scalar(text("select count(*) from operation_buyers")) == 0


def test_preview_and_atomic_apply_roundtrip(
    sales_ops_client: TestClient, project_id: str, buyer_id: str, db: Session
) -> None:
    second = sales_ops_client.post(
        f"/api/v1/projects/{project_id}/sales/clients", json={"display_name": "Second buyer"}
    )
    assert second.status_code == 201, second.text
    book = workbook(sales_ops_client, project_id)
    sheet = book["Clients"]
    for row in (3, 4):
        sheet.cell(row, 6, "Golden Visa")
        sheet.cell(row, 7, "Yes")
        sheet.cell(row, 8, date(2026, 10, 29))
    preview = upload(sales_ops_client, project_id, book)
    assert preview.status_code == 200, preview.text
    assert preview.json()["changed_clients"] == 2
    assert preview.json()["change_count"] == 4
    assert db.scalar(text("select count(*) from operation_buyers")) == 0
    result = upload(sales_ops_client, project_id, book, "apply")
    assert result.status_code == 200, result.text
    assert result.json()["applied"]
    data = read(sales_ops_client, project_id)
    assert data["summaries"][0]["yes"] == 2
    assert data["summaries"][0]["completion_percent"] == 100
    assert all(b["milestones"][0]["completed_date"] == "2026-10-29" for b in data["buyers"])
    assert (
        db.scalar(
            text("select count(*) from audit_events where action = 'operations.buyer_updated'")
        )
        == 2
    )
    replay = upload(sales_ops_client, project_id, book, "apply")
    assert replay.status_code == 409  # initial default pipeline materialized
    fresh = workbook(sales_ops_client, project_id)
    assert upload(sales_ops_client, project_id, fresh).json()["changed_clients"] == 0


@pytest.mark.parametrize("case", ["duplicate", "foreign", "date", "formula", "status", "visa"])
def test_invalid_file_saves_nothing(
    sales_ops_client: TestClient, project_id: str, buyer_id: str, db: Session, case: str
) -> None:
    book = workbook(sales_ops_client, project_id)
    sheet = book["Clients"]
    sheet["G3"] = "Yes"
    if case == "duplicate":
        sheet.append([c.value for c in sheet[3]])
    if case == "foreign":
        sheet["A3"] = str(uuid.uuid4())
    if case == "date":
        sheet["H3"] = "29/10/2026"
    if case == "formula":
        sheet["H3"] = "=TODAY()"
    if case == "status":
        sheet["I3"] = "Maybe"
    if case == "visa":
        sheet["U3"] = "Yes"
    response = upload(sales_ops_client, project_id, book, "apply")
    assert response.status_code == 422 or response.json()["issues"], response.text
    assert db.scalar(text("select count(*) from operation_buyers")) == 0
    assert db.scalar(text("select count(*) from operation_progress")) == 0


def test_stale_progress_and_pipeline(
    sales_ops_client: TestClient, project_id: str, buyer_id: str
) -> None:
    data = read(sales_ops_client, project_id)
    assert (
        sales_ops_client.put(
            f"{url(project_id)}/buyers/{buyer_id}", json=payload(data, buyer_id)
        ).status_code
        == 204
    )
    book = workbook(sales_ops_client, project_id)
    book["Clients"]["G3"] = "Yes"
    data = read(sales_ops_client, project_id)
    assert (
        sales_ops_client.put(
            f"{url(project_id)}/buyers/{buyer_id}", json=payload(data, buyer_id)
        ).status_code
        == 204
    )
    report = upload(sales_ops_client, project_id, book, "apply").json()
    assert report["issues"] and not report["applied"]
    book = workbook(sales_ops_client, project_id)
    data = read(sales_ops_client, project_id)
    assert (
        sales_ops_client.put(f"{url(project_id)}/pipeline", json=stage_payload(data)).status_code
        == 204
    )
    assert upload(sales_ops_client, project_id, book, "apply").status_code == 409


def test_blank_preserves_explicit_clear_removes(
    sales_ops_client: TestClient, project_id: str, buyer_id: str
) -> None:
    book = workbook(sales_ops_client, project_id)
    book["Clients"]["G3"] = "Yes"
    book["Clients"]["H3"] = "2026-10-29"
    assert upload(sales_ops_client, project_id, book, "apply").json()["applied"]
    book = workbook(sales_ops_client, project_id)
    book["Clients"]["G3"] = None
    book["Clients"]["H3"] = None
    assert upload(sales_ops_client, project_id, book).json()["change_count"] == 0
    book["Clients"]["H3"] = "CLEAR"
    assert upload(sales_ops_client, project_id, book, "apply").json()["applied"]
    assert (
        read(sales_ops_client, project_id)["buyers"][0]["milestones"][0]["completed_date"] is None
    )


def test_import_denies_reader_and_wrong_scope(
    sales_ops_client: TestClient,
    finance_client: TestClient,
    admin_client: TestClient,
    project_id: str,
    buyer_id: str,
    db: Session,
) -> None:
    from tests.factories import client_for, make_user
    from tests.modules.conftest import grant_access

    book = workbook(sales_ops_client, project_id)
    assert upload(finance_client, project_id, book).status_code == 403
    assert upload(sales_ops_client, str(uuid.uuid4()), book).status_code == 404
    user = make_user(db, email="phase-import@example.com", roles=("sales_operations",))
    grant_access(admin_client, project_id, user)
    assert (
        admin_client.patch(
            f"/api/v1/projects/{project_id}/access/{user.id}/phase-scope",
            json={"phase_scope": "selected"},
        ).status_code
        == 200
    )
    with client_for(user.email) as restricted:
        assert restricted.get(f"{url(project_id)}/import/template").status_code == 404
        assert upload(restricted, project_id, book).status_code == 404


def test_import_cannot_overwrite_sales_signature(
    sales_ops_client: TestClient,
    project_id: str,
    submitted_sale: str,
) -> None:
    book = workbook(sales_ops_client, project_id)
    book["Clients"]["M3"] = "Yes"
    result = upload(sales_ops_client, project_id, book, "apply")
    assert result.status_code == 200, result.text
    assert "linked to Sales" in result.json()["issues"][0]["message"]
    assert not result.json()["applied"]


def test_apply_rolls_back_prior_rows_if_later_write_fails(
    sales_ops_client: TestClient,
    project_id: str,
    buyer_id: str,
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.core.errors import ConflictError
    from app.modules.sales import operations

    second = sales_ops_client.post(
        f"/api/v1/projects/{project_id}/sales/clients", json={"display_name": "Second"}
    )
    assert second.status_code == 201
    book = workbook(sales_ops_client, project_id)
    book["Clients"]["G3"] = "Yes"
    book["Clients"]["G4"] = "Yes"
    original = operations.apply_buyer_update
    calls = 0

    def fail_second(*args: object, **kwargs: object) -> None:
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ConflictError("Synthetic conflict")
        original(*args, **kwargs)

    monkeypatch.setattr(operations, "apply_buyer_update", fail_second)
    assert upload(sales_ops_client, project_id, book, "apply").status_code == 409
    assert db.scalar(text("select count(*) from operation_buyers")) == 0
    assert (
        db.scalar(
            text("select count(*) from audit_events where action = 'operations.buyer_updated'")
        )
        == 0
    )
