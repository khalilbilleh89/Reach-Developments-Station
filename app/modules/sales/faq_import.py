"""Bounded, project-bound FAQ spreadsheets; validate before one atomic apply."""

import io
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError as SchemaError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ValidationError
from app.core.xlsx import check_archive
from app.modules.access.dependencies import ActorContext
from app.modules.projects.models import Project
from app.modules.projects.service import lock_project
from app.modules.sales import faq
from app.modules.sales.faq_models import CommercialFaq
from app.modules.sales.faq_schemas import FaqInput
from app.modules.sales.schemas import SalesImportIssueRead, SalesImportReportRead

MAX_ROWS = 1000


def template(session: Session, project: Project, actor: ActorContext) -> bytes:
    faq.require_writer(session, project, actor)
    book = Workbook()
    sheet = book.active
    sheet.title = "FAQs"
    sheet.append(["Question", "Answer"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{MAX_ROWS + 1}"
    sheet.column_dimensions["A"].width = 55
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=2):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions = book.create_sheet("Instructions")
    for row in [
        ("Template version", "commercial_faqs_v1"),
        ("Project", str(project.id)),
        ("Fill FAQs", "One question and answer per row, starting at row 2. Keep both headers."),
        ("Limits", "Up to 1,000 FAQs and 5 MB. Question: 500 characters; answer: 20,000."),
        (
            "Text",
            "Both fields are required. Paste plain text, never formulas. "
            "Line breaks and Unicode are preserved.",
        ),
        (
            "Import",
            "Validate, fix any reported rows, then import. Any issue prevents the entire batch.",
        ),
        (
            "Duplicates",
            "Repeated questions, ignoring case and surrounding spaces, are rejected. "
            "Existing FAQs are never overwritten.",
        ),
        (
            "Blank rows",
            "Completely blank rows are ignored. "
            "Use Edit in the FAQ library to change an existing answer.",
        ),
    ]:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 100
    for row in instructions:
        instructions.row_dimensions[row[0].row].height = 44
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for page in book:
        for cell in page[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173F42")
        page.sheet_view.showGridLines = False
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def rows(body: bytes, project_id: str) -> list[tuple[int, dict[str, str]]]:
    check_archive(body)
    try:
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
        try:
            if not {"Instructions", "FAQs"}.issubset(book.sheetnames):
                raise ValidationError("Download the FAQ template and keep its sheet names.")
            if book["Instructions"]["B1"].value != "commercial_faqs_v1":
                raise ValidationError("Download the current FAQ template.")
            if book["Instructions"]["B2"].value != project_id:
                raise ValidationError(
                    "This template belongs to another project. Download this project's template."
                )
            sheet = book["FAQs"]
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            header = [str(cell.value or "").strip() for cell in next(iterator, ())]
            while header and not header[-1]:
                header.pop()
            if sorted(header) != ["Answer", "Question"]:
                raise ValidationError("Keep the Question and Answer headers exactly once each.")
            parsed = []
            for number, cells in enumerate(iterator, 2):
                if number > MAX_ROWS + 1001:
                    raise ValidationError("Too many worksheet rows. Remove excess blank rows.")
                values = {}
                for position, cell in enumerate(cells):
                    if cell.data_type in {"f", "e"} or (
                        isinstance(cell.value, str) and cell.value.startswith("=")
                    ):
                        raise ValidationError(
                            f"Row {number} contains a formula or cell error. "
                            "Paste plain text instead."
                        )
                    value = "" if cell.value is None else str(cell.value).strip()
                    if position >= len(header):
                        if value:
                            raise ValidationError(
                                f"Row {number} has data outside the template columns."
                            )
                    else:
                        values[header[position].lower()] = value
                if not any(values.values()):
                    continue
                parsed.append((number, values))
                if len(parsed) > MAX_ROWS:
                    raise ValidationError("Import at most 1,000 FAQs at a time.")
            if not parsed:
                raise ValidationError("No FAQs found. Fill the FAQs sheet starting at row 2.")
            return parsed
        finally:
            book.close()
    except (KeyError, ValueError, TypeError, OSError, ParseError, zipfile.BadZipFile) as exc:
        raise ValidationError("That file is not a readable FAQ template.") from exc


def process(
    session: Session, project: Project, actor: ActorContext, body: bytes, *, apply: bool
) -> SalesImportReportRead:
    faq.require_writer(session, project, actor)
    parsed = rows(body, str(project.id))
    lock_project(session, project.id)
    seen = {
        question.strip().casefold()
        for question in session.scalars(
            select(CommercialFaq.question).where(CommercialFaq.project_id == project.id)
        )
    }
    issues = []
    payloads = []
    for number, values in parsed:
        try:
            payload = FaqInput.model_validate(values)
        except SchemaError as exc:
            issues.extend(
                SalesImportIssueRead(row=number, column=str(error["loc"][0]), message=error["msg"])
                for error in exc.errors(include_input=False, include_url=False)
            )
            continue
        key = payload.question.casefold()
        if key in seen:
            issues.append(
                SalesImportIssueRead(
                    row=number,
                    column="question",
                    message="This question already exists or is repeated in this file. "
                    "Remove the duplicate row.",
                )
            )
        seen.add(key)
        payloads.append(payload)
    applied = apply and not issues
    try:
        if applied:
            for payload in payloads:
                faq.create(session, project, actor, payload, commit=False)
            session.commit()
        else:
            session.rollback()
    except Exception:
        session.rollback()
        raise
    return SalesImportReportRead(
        total_rows=len(parsed),
        created_count=len(payloads) if applied else 0,
        applied=applied,
        error_count=len(issues),
        issues=issues[:200],
        issues_truncated=len(issues) > 200,
    )
