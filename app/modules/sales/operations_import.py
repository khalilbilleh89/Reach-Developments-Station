"""Project-specific Operations workbook: bounded input, preview, atomic update."""

import io
import zipfile
from datetime import date, datetime
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session

from app.core.errors import ConflictError, ServiceError, ValidationError
from app.core.xlsx import check_archive
from app.modules.access.dependencies import ActorContext
from app.modules.projects.models import Project
from app.modules.sales import operations
from app.modules.sales.operations_schemas import BuyerInput, OperationsRead, ProgressInput

MAX_ROWS = 5000
FIXED = [
    "Client ID",
    "Version",
    "Client reference",
    "Client name",
    "Purchases",
    "Purpose of Purchase",
]
PURPOSES = {
    "Investment Only": "investment_only",
    "Golden Visa": "golden_visa",
    "Not recorded": None,
}
STATUSES = {"Yes": True, "No": False, "Not recorded": None}


def _headers(data: OperationsRead) -> list[str]:
    return FIXED + [
        f"{stage.id}:{field}"
        for stage in data.stages
        if stage.is_active
        for field in ("status", "date")
    ]


def template(session: Session, project: Project, actor: ActorContext) -> bytes:
    operations.require_access(session, project, actor)
    data = operations.read(session, project, actor)
    if len(data.buyers) > MAX_ROWS:
        raise ValidationError(f"The Excel workflow supports up to {MAX_ROWS} clients per project.")
    book = Workbook()
    notes = book.active
    notes.title = "Instructions"
    for row in [
        ["Template", "operations_v1"],
        ["Project", str(project.id)],
        ["Pipeline version", data.pipeline_version],
        [
            "Workflow",
            "Edit Clients, validate the file, review the changes, then import. All "
            "changes save together or none do.",
        ],
        [
            "Input",
            "Blue cells are editable. Grey cells are reference or linked Sales "
            "evidence. Keep both header rows and client IDs.",
        ],
        [
            "Blank cells",
            "Blank means keep the current value. Choose Not recorded to clear a "
            "status or purpose. Type CLEAR to remove a date.",
        ],
        [
            "Dates",
            "Use Excel dates or YYYY-MM-DD text. A date requires Yes. No / Not "
            "recorded clears the date; clear an old date cell when changing status.",
        ],
        [
            "Golden Visa",
            "Choose Golden Visa before entering visa milestones. Existing visa "
            "evidence is preserved when purpose changes.",
        ],
        [
            "Sales",
            "SPA with current sales is read-only. Correct buyer signatures in Sales, "
            "then download a fresh template.",
        ],
        [
            "Freshness",
            "If client progress or project stages change after download, download a "
            "new template. Do not edit the hidden ID/version columns.",
        ],
        [
            "Limits",
            f"{MAX_ROWS} clients; 5 MB. Values only: no formulas, macros or cell "
            "errors. No clients are created or deleted.",
        ],
    ]:
        notes.append(row)
    sheet = book.create_sheet("Clients")
    sheet.append(_headers(data))
    sheet.row_dimensions[1].hidden = True
    stages = [stage for stage in data.stages if stage.is_active]
    sheet.append(
        FIXED
        + [
            f"{'Purchase' if stage.section == 'property_purchase' else 'Golden Visa'}"
            f" · {stage.label} · {field}"
            for stage in stages
            for field in ("Status", "Date")
        ]
    )
    for buyer in data.buyers:
        purpose = next(label for label, value in PURPOSES.items() if value == buyer.purpose)
        values = [
            str(buyer.id),
            buyer.version,
            buyer.number,
            buyer.name,
            "; ".join(p.number for p in buyer.purchases),
            purpose,
        ]
        for stage in stages:
            milestone = next(m for m in buyer.milestones if m.stage_id == stage.id)
            values += [
                next(label for label, value in STATUSES.items() if value is milestone.completed),
                milestone.completed_date,
            ]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Buyer text must never become a formula.
            if cell.column < 6:
                cell.font = Font(color="526369")
                cell.fill = PatternFill("solid", fgColor="F5F7FB")
        for column in range(6, sheet.max_column + 1):
            cell = sheet.cell(sheet.max_row, column)
            linked = (
                column > 6
                and next(
                    m for m in buyer.milestones if m.stage_id == stages[(column - 7) // 2].id
                ).total_sales
                > 0
            )
            cell.protection = Protection(locked=linked)
            cell.font = Font(color="526369" if linked else "245CCE")
            cell.fill = PatternFill("solid", fgColor="F5F7FB" if linked else "FFFFFF")
            if column > 6 and column % 2 == 0:
                cell.number_format = "yyyy-mm-dd"
    sheet.column_dimensions["A"].hidden = True
    sheet.column_dimensions["B"].hidden = True
    sheet.freeze_panes = "G3"
    sheet.auto_filter.ref = f"C2:{get_column_letter(sheet.max_column)}{max(2, sheet.max_row)}"
    for choices, columns in [
        (list(PURPOSES), [6]),
        (list(STATUSES), list(range(7, sheet.max_column + 1, 2))),
    ]:
        validation = DataValidation(
            type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True
        )
        validation.errorTitle = "Choose a listed value"
        validation.error = "Use the dropdown or leave blank to keep the current value."
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        for column in columns:
            validation.add(
                f"{get_column_letter(column)}3:{get_column_letter(column)}{max(3, sheet.max_row)}"
            )
    for page in book:
        page.sheet_view.showGridLines = False
        header_row = 2 if page == sheet else 1
        for cell in page[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245CCE")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for column in range(1, page.max_column + 1):
            page.column_dimensions[get_column_letter(column)].width = (
                28 if page == sheet else (24 if column == 1 else 105)
            )
    sheet.row_dimensions[2].height = 64
    sheet.column_dimensions["D"].width = 34
    for row in notes:
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 42
    book.active = 1
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def _rows(body: bytes, project: Project, data: OperationsRead) -> list[tuple[int, dict]]:
    check_archive(body)
    try:
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
        try:
            if not {"Instructions", "Clients"}.issubset(book.sheetnames):
                raise ValidationError("Use the Operations template downloaded from this project.")
            notes = book["Instructions"]
            if notes["B1"].value != "operations_v1" or notes["B2"].value != str(project.id):
                raise ValidationError("This is not this project's Operations template.")
            if notes["B3"].value != data.pipeline_version:
                raise ConflictError("The pipeline changed. Download a fresh template.")
            sheet = book["Clients"]
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(max_col=len(_headers(data)) + 1)
            header = [c.value for c in next(iterator, ())]
            while header and header[-1] is None:
                header.pop()
            if header != _headers(data):
                raise ValidationError(
                    "The stage columns changed. Download a fresh template and keep the headers."
                )
            next(iterator, ())  # human-readable headings
            result = []
            for number, cells in enumerate(iterator, 3):
                if number > MAX_ROWS + 2:
                    raise ValidationError(
                        f"Use at most {MAX_ROWS} client rows; remove excess blank rows."
                    )
                if not any(c.value is not None for c in cells):
                    continue
                if any(c.data_type in {"f", "e"} for c in cells):
                    raise ValidationError(
                        f"Row {number}: formulas and cell errors are not accepted. Paste values."
                    )
                if len(cells) > len(header) and cells[-1].value is not None:
                    raise ValidationError(f"Row {number}: data outside the template columns.")
                result.append(
                    (number, {key: cells[index].value for index, key in enumerate(header)})
                )
            if not result:
                raise ValidationError("There are no client rows in this file.")
            return result
        finally:
            book.close()
    except (KeyError, ValueError, TypeError, OSError, ParseError, zipfile.BadZipFile) as exc:
        raise ValidationError("That file is not a readable Operations workbook.") from exc


def _date(value: object) -> date | None:
    if value == "CLEAR":
        return None
    if isinstance(value, datetime):
        if value.time().isoformat() != "00:00:00":
            raise ValidationError("Use a date without a time.")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and len(value) == 10:
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    raise ValidationError("Use an Excel date, YYYY-MM-DD, or CLEAR.")


def _display(completed: bool | None, day: date | None) -> str:
    status = next(k for k, v in STATUSES.items() if v is completed)
    return f"{status} · {day or 'No date'}"


def process(
    session: Session,
    project: Project,
    actor: ActorContext,
    body: bytes,
    *,
    apply: bool,
    reason: str,
) -> dict:
    operations._write(session, project, actor)
    data = operations.read(session, project, actor)
    rows = _rows(body, project, data)
    buyers = {str(b.id): b for b in data.buyers}
    seen = set()
    updates = []
    changes = []
    issues = []
    for number, values in rows:
        try:
            key = str(values["Client ID"])
            if key in seen:
                raise ValidationError("Duplicate client row; keep one row per client.")
            seen.add(key)
            buyer = buyers.get(key)
            if buyer is None:
                raise ValidationError("Client not found in this project's accessible register.")
            if values["Version"] != buyer.version:
                raise ConflictError("Client progress changed. Download a fresh template.")
            purpose_value = values["Purpose of Purchase"]
            if purpose_value is not None and purpose_value not in PURPOSES:
                raise ValidationError(
                    "Purpose must be Investment Only, Golden Visa or Not recorded."
                )
            purpose = buyer.purpose if purpose_value is None else PURPOSES[purpose_value]
            progress = []
            row_changes = []
            if purpose != buyer.purpose:
                row_changes.append(
                    {
                        "row": number,
                        "client": buyer.name,
                        "field": "Purpose of Purchase",
                        "before": buyer.purpose or "Not recorded",
                        "after": purpose or "Not recorded",
                    }
                )
            for stage in data.stages:
                if not stage.is_active:
                    continue
                old = next(m for m in buyer.milestones if m.stage_id == stage.id)
                status = values[f"{stage.id}:status"]
                raw_date = values[f"{stage.id}:date"]
                if status is not None and status not in STATUSES:
                    raise ValidationError(f"{stage.label}: use Yes, No or Not recorded.")
                completed = old.completed if status is None else STATUSES[status]
                completed_date = old.completed_date if raw_date is None else _date(raw_date)
                if completed is not True:
                    if raw_date not in (None, "CLEAR"):
                        raise ValidationError(
                            f"{stage.label}: a date requires Yes; clear the date cell."
                        )
                    completed_date = None
                if (completed, completed_date) == (old.completed, old.completed_date):
                    continue
                if old.total_sales:
                    raise ConflictError(
                        f"{stage.label} is linked to Sales. Correct the buyer signature in Sales."
                    )
                if stage.section == "golden_visa" and purpose != "golden_visa":
                    raise ValidationError(
                        f"{stage.label}: choose Golden Visa before entering visa progress."
                    )
                progress.append(
                    ProgressInput(
                        stage_id=stage.id, completed=completed, completed_date=completed_date
                    )
                )
                row_changes.append(
                    {
                        "row": number,
                        "client": buyer.name,
                        "field": stage.label,
                        "before": _display(old.completed, old.completed_date),
                        "after": _display(completed, completed_date),
                    }
                )
            if progress or purpose != buyer.purpose:
                updates.append(
                    (
                        buyer,
                        BuyerInput(
                            expected_version=buyer.version,
                            pipeline_version=data.pipeline_version,
                            purpose=purpose,
                            progress=progress,
                            reason=reason or "Excel preview",
                        ),
                    )
                )
                changes.extend(row_changes)
        except ServiceError as exc:
            issues.append({"row": number, "message": str(exc)})
    report = {
        "total_rows": len(rows),
        "changed_clients": len(updates),
        "unchanged_clients": len(rows) - len(updates),
        "change_count": len(changes),
        "changes": changes,
        "issues": issues,
        "applied": False,
    }
    if issues or not apply:
        session.rollback()
        return report
    if not reason.strip():
        raise ValidationError("Give a reason for this import.")
    try:
        for buyer, payload in updates:
            operations.apply_buyer_update(session, project, actor, data, buyer, payload)
        session.commit()
    except Exception:
        session.rollback()
        raise
    report["applied"] = True
    return report
