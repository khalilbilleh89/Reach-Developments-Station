"""Excel input for the existing Buyer and Agent creation contracts.

Validation stages the same domain writes inside savepoints, then rolls them all
back. Apply checks the file again under the project lock and commits once.
"""

import io
import zipfile
from typing import Literal
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import ValidationError as SchemaError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.errors import ServiceError, ValidationError
from app.core.xlsx import check_archive
from app.modules.access.dependencies import ActorContext
from app.modules.projects.models import Project
from app.modules.projects.service import lock_project
from app.modules.sales import agents, permissions, service
from app.modules.sales.schemas import ClientCreateRequest, SalesAgentCreateRequest

Kind = Literal["buyers", "agents"]
MAX_ROWS = 1000
COLUMNS = {
    "buyers": ("display_name", "email", "phone", "sole_purchaser", "agent_id"),
    "agents": ("display_name", "country", "branch", "branch_leader"),
}


def template(session: Session, project: Project, actor: ActorContext, kind: Kind) -> bytes:
    permissions.require_client_writer(actor)
    book = Workbook()
    instructions = book.active
    instructions.title = "Instructions"
    instructions.append(["Template version", f"sales_{kind}_v1"])
    instructions.append(["Project", str(project.id)])
    instructions.append(
        [
            "How to use",
            f"Fill {kind.title()} starting at row 2. Keep the headers. Maximum {MAX_ROWS} records.",
        ]
    )
    instructions.append(
        ["Required", "display_name is required. Other fields are optional. Blank rows are ignored."]
    )
    instructions.append(
        [
            "Import",
            (
                "Validate first. Any error prevents the entire batch. Existing records "
                "are never updated."
            ),
        ]
    )
    instructions.append(
        [
            "Text",
            (
                "Keep phone numbers and identifiers formatted as Text to preserve "
                "leading zeros. Paste values, never formulas."
            ),
        ]
    )
    if kind == "buyers":
        instructions.append(
            [
                "sole_purchaser",
                (
                    "TRUE or FALSE; blank means TRUE, matching Register buyer. TRUE records "
                    "100% ownership in the buyer name. For joint buyers, use FALSE and add "
                    "parties after import."
                ),
            ]
        )
        instructions.append(
            [
                "agent_id",
                (
                    "Optional: copy an ID from Agent reference. Country, branch and branch "
                    "leader come from that Agent. Import Agents first if needed."
                ),
            ]
        )
        reference = book.create_sheet("Agent reference")
        reference.append(["agent_id", "display_name", "country", "branch", "branch_leader"])
        for agent in agents.list_agents(session, project=project, actor=actor, active_only=True):
            reference.append(
                [
                    str(agent.id),
                    agent.display_name,
                    agent.country,
                    agent.branch,
                    agent.branch_leader,
                ]
            )
            for cell in reference[reference.max_row]:
                cell.data_type = "s"
    sheet = book.create_sheet(kind.title())
    sheet.append(COLUMNS[kind])
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=len(COLUMNS[kind])):
        for cell in row:
            cell.number_format = "@"
    for page in book:
        page.freeze_panes = "A2"
        for cell in page[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173F42")
        for index in range(1, page.max_column + 1):
            page.column_dimensions[get_column_letter(index)].width = (
                28 if page != instructions else (24 if index == 1 else 110)
            )
    for row in instructions:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        instructions.row_dimensions[row[0].row].height = 45
    if kind == "buyers":
        book["Agent reference"].column_dimensions["A"].width = 40
        sheet.column_dimensions["E"].width = 40
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def rows(body: bytes, kind: Kind, project_id: str) -> list[tuple[int, dict[str, str]]]:
    check_archive(body)
    try:
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
        try:
            if not {"Instructions", kind.title()}.issubset(book.sheetnames):
                raise ValidationError("Download the template for this tab and fill that in.")
            if book["Instructions"]["B1"].value != f"sales_{kind}_v1":
                raise ValidationError("The template version does not match this tab.")
            if book["Instructions"]["B2"].value != project_id:
                raise ValidationError(
                    "This template belongs to another project. Download this project's template."
                )
            sheet = book[kind.title()]
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            headers = [str(cell.value or "").strip() for cell in next(iterator, ())]
            while headers and not headers[-1]:
                headers.pop()
            if sorted(headers) != sorted(COLUMNS[kind]):
                raise ValidationError(
                    "The column headers do not match the template. Keep each header exactly once."
                )
            result = []
            for number, cells in enumerate(iterator, 2):
                if number > MAX_ROWS + 1001:
                    raise ValidationError("Too many worksheet rows. Remove excess blank rows.")
                values = {}
                for position, cell in enumerate(cells):
                    if cell.data_type in {"f", "e"} or (
                        isinstance(cell.value, str) and cell.value.startswith("=")
                    ):
                        raise ValidationError(
                            f"Row {number} contains a formula or cell error. Paste values instead."
                        )
                    text = "" if cell.value is None else str(cell.value).strip()
                    if position >= len(headers):
                        if text:
                            raise ValidationError(
                                f"Row {number} has data outside the template columns."
                            )
                    else:
                        values[headers[position]] = text
                if not any(values.values()):
                    continue
                result.append((number, values))
                if len(result) > MAX_ROWS:
                    raise ValidationError(f"Import at most {MAX_ROWS} records at a time.")
            if not result:
                raise ValidationError(
                    "The workbook contains no records. Fill the template starting at row 2."
                )
            return result
        finally:
            book.close()
    except (KeyError, ValueError, TypeError, OSError, ParseError, zipfile.BadZipFile) as exc:
        raise ValidationError("That file is not a readable Excel template.") from exc


def payload(values: dict[str, str], kind: Kind) -> dict[str, object]:
    fields = {key: value for key, value in values.items() if value}
    if kind == "agents":
        return SalesAgentCreateRequest.model_validate(fields).model_dump()
    sole = str(fields.pop("sole_purchaser", "true")).casefold()
    if sole not in {"true", "false"}:
        raise ValidationError("sole_purchaser must be TRUE or FALSE.")
    if sole == "true":
        fields["sole_purchaser_name"] = fields.get("display_name", "")
    return ClientCreateRequest.model_validate(fields).model_dump(exclude_unset=True)


def process(
    session: Session, project: Project, actor: ActorContext, kind: Kind, body: bytes, *, apply: bool
) -> dict[str, object]:
    permissions.require_client_writer(actor)
    project = lock_project(session, project.id)
    if kind == "buyers":
        permissions.require_operational_project(project)
    parsed = rows(body, kind, str(project.id))
    # Respect advisor visibility: duplicate checks cannot disclose other buyers.
    existing = (
        service.list_clients(session, project=project, actor=actor)
        if kind == "buyers"
        else agents.list_agents(session, project=project, actor=actor)
    )
    identity_fields = ("display_name", "email", "phone") if kind == "buyers" else COLUMNS[kind]

    def identity(fields: dict[str, object]) -> tuple[str, ...]:
        return tuple(str(fields.get(key) or "").strip().casefold() for key in identity_fields)

    seen = {identity({key: getattr(record, key) for key in identity_fields}) for record in existing}
    issues = []
    for number, values in parsed:
        try:
            fields = payload(values, kind)
            key = identity(fields)
            if key in seen:
                raise ValidationError(
                    "A matching record already exists or is repeated in this file. Remove "
                    "the duplicate row."
                )
            with session.begin_nested():
                if kind == "buyers":
                    service.create_client(
                        session, project=project, actor=actor, commit=False, **fields
                    )
                else:
                    agents.create_agent(
                        session, project=project, actor=actor, fields=fields, commit=False
                    )
            seen.add(key)
        except SchemaError as exc:
            for error in exc.errors(include_input=False, include_url=False):
                issues.append(
                    {
                        "row": number,
                        "column": ".".join(map(str, error["loc"])),
                        "message": error["msg"],
                    }
                )
        except ServiceError as exc:
            issues.append({"row": number, "column": None, "message": exc.detail})
        except IntegrityError:
            issues.append(
                {
                    "row": number,
                    "column": None,
                    "message": "This row conflicts with an existing record.",
                }
            )
    applied = apply and not issues
    if applied:
        session.commit()
    else:
        session.rollback()
    return {
        "total_rows": len(parsed),
        "created_count": len(parsed) if applied else 0,
        "applied": applied,
        "error_count": len(issues),
        "issues": issues[:200],
        "issues_truncated": len(issues) > 200,
    }
